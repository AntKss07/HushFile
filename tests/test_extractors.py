"""Tests for file content extractors."""

import csv
from pathlib import Path

import docx
import fitz
import openpyxl
import pytest

from butler.config import ButlerConfig
from butler.extractors import ExtractionResult, extract_content


@pytest.fixture
def config():
    return ButlerConfig()


def test_text_extractor(tmp_path: Path, config: ButlerConfig):
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Meeting notes from project planning", encoding="utf-8")
    res = extract_content(txt_file, config)
    assert isinstance(res, ExtractionResult)
    assert "Meeting notes" in res.text
    assert res.metadata["extension"] == ".txt"


def test_code_extractor(tmp_path: Path, config: ButlerConfig):
    py_file = tmp_path / "script.py"
    py_file.write_text("def hello_world():\n    return 'OK'\n", encoding="utf-8")
    res = extract_content(py_file, config)
    assert "def hello_world" in res.text
    assert res.metadata["extension"] == ".py"


def test_pdf_extractor(tmp_path: Path, config: ButlerConfig):
    pdf_file = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Quarterly Financial Statement 2026\nRevenue: $500,000")
    doc.save(str(pdf_file))
    doc.close()

    res = extract_content(pdf_file, config)
    assert "Quarterly Financial Statement" in res.text
    assert "Revenue: $500,000" in res.text
    assert res.metadata["page_count"] == 1


def test_docx_extractor(tmp_path: Path, config: ButlerConfig):
    doc_file = tmp_path / "document.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph of executive summary")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Price"
    doc.save(str(doc_file))

    res = extract_content(doc_file, config)
    assert "First paragraph of executive summary" in res.text
    assert "Item | Price" in res.text


def test_xlsx_extractor(tmp_path: Path, config: ButlerConfig):
    xlsx_file = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Amount"])
    ws.append(["January", 1500])
    wb.save(str(xlsx_file))

    res = extract_content(xlsx_file, config)
    assert "Sheet: Sales" in res.text
    assert "Month | Amount" in res.text
    assert "January | 1500" in res.text


def test_csv_extractor(tmp_path: Path, config: ButlerConfig):
    csv_file = tmp_path / "items.csv"
    with open(csv_file, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Name", "Cost"])
        writer.writerow(["1", "Widget", "9.99"])

    res = extract_content(csv_file, config)
    assert "ID | Name | Cost" in res.text
    assert "1 | Widget | 9.99" in res.text


def test_binary_archive_extractor(tmp_path: Path, config: ButlerConfig):
    zip_file = tmp_path / "archive.zip"
    zip_file.write_bytes(b"PK\x03\x04testdummydata")
    res = extract_content(zip_file, config)
    assert res.text == ""
    assert res.metadata["is_binary_package"] is True
    assert res.metadata["extension"] == ".zip"


def test_corrupt_or_missing_file_graceful(tmp_path: Path, config: ButlerConfig):
    missing_file = tmp_path / "does_not_exist.pdf"
    res = extract_content(missing_file, config)
    assert isinstance(res, ExtractionResult)
    assert res.text == ""
    assert res.warning is not None
