"""Generate synthetic dataset and labels for HushFile benchmark."""

import csv
from pathlib import Path

import docx
import fitz  # PyMuPDF
from PIL import Image, ImageDraw


def generate_synthetic_dataset(output_dir: Path | str | None = None) -> Path:
    """Create sample files and labels.csv across core categories."""
    base = Path(output_dir) if output_dir else Path(__file__).parent / "dataset"
    base.mkdir(parents=True, exist_ok=True)
    labels_file = Path(__file__).parent / "labels.csv"

    samples = [
        # Finance
        ("invoice_cloud_server_2026.pdf", "Finance", "pdf", "Invoice INV-90412\nDue Date: 2026-11-01\nVendor: CloudHosting Inc.\nAmount Due: $149.00\nServices: Dedicated GPU Compute Node"),
        ("quarterly_tax_estimate.pdf", "Finance", "pdf", "Estimated Tax Declaration 2026\nQuarter 3\nTotal Federal Income Tax Withheld: $4,200.00"),
        ("bank_monthly_statement.pdf", "Finance", "pdf", "First National Bank\nAccount Statement Period: September 2026\nOpening Balance: $12,450.00\nClosing Balance: $14,120.00"),
        # Receipts
        ("supermarket_grocery_receipt.pdf", "Receipts", "pdf", "Store #104 Organic Market\nTotal Items: 6\nSubtotal: $45.20\nTax: $3.80\nTotal Paid: $49.00\nPayment Method: Visa Ending in 4102"),
        ("coffee_shop_bill.txt", "Receipts", "txt", "Cafe Central\n1x Americano $4.50\n1x Croissant $3.75\nTotal: $8.25\nThank you for visiting!"),
        # Work
        ("project_proposal_draft.docx", "Work", "docx", "Project Apollo: Next-Generation Agent Infrastructure.\nExecutive Summary: This document details the architectural roadmap and deliverables for Q4."),
        ("client_nda_signed.pdf", "Work", "pdf", "MUTUAL NONDISCLOSURE AGREEMENT\nBetween Acme Corp and Client. Confidential information shall remain protected for 5 years."),
        # Education
        ("biology_lecture_04_mitosis.pdf", "Education", "pdf", "Cellular Biology 101 - Lecture 4: Mitosis and Meiosis\nProfessor Dr. Smith\nSyllabus overview and exam schedule."),
        ("calculus_homework_set3.pdf", "Education", "pdf", "Calculus II - Homework Assignment 3\nSolve the following differential equations with initial boundary conditions."),
        # Health
        ("annual_bloodwork_panel.pdf", "Health", "pdf", "Medical Laboratory Diagnostic Report\nPatient: John Doe\nTest: Comprehensive Metabolic Panel\nGlucose: 92 mg/dL (Normal)\nHemoglobin: 14.5 g/dL"),
        ("prescription_pharmacy.pdf", "Health", "pdf", "Rx Pharmacy Order\nPrescription details: Amoxicillin 500mg, twice daily for 10 days."),
        # Travel
        ("flight_boarding_pass_nyc.pdf", "Travel", "pdf", "AIRLINE BOARDING PASS\nPassenger: Traveler/Jane\nFlight: PA402 JFK to SFO\nDate: 2026-11-12 Gate: B22 Seat: 14A"),
        ("hotel_reservation_confirmation.pdf", "Travel", "pdf", "Grand Hotel Booking Confirmation\nCheck-in: Nov 12, 2026\nCheck-out: Nov 16, 2026\nTotal Stay: $840.00"),
        # Identity
        ("passport_copy_scan.pdf", "Identity", "pdf", "REPUBLIC PASSPORT\nSurname: TRAVELER\nGiven Names: JANE\nNationality: USA\nDate of Birth: 1995-04-12"),
        # Code
        ("data_pipeline.py", "Code", "code", '"""ETL pipeline script for aggregating nightly telemetry."""\nimport json\nimport sys\ndef process_records(stream):\n    return [json.loads(line) for line in stream]\n'),
        ("docker_compose.yaml", "Code", "code", "version: '3.8'\nservices:\n  web:\n    image: nginx:alpine\n    ports:\n      - '8080:80'\n"),
        # Screenshots
        ("screenshot_system_settings.png", "Screenshots", "img", "Screenshot of Desktop display settings and sound options"),
        # Archives
        ("backup_logs_2026.zip", "Archives", "binary", b"PK\x03\x04zipdummy"),
        # Installers
        ("git_installer_x64.exe", "Installers", "binary", b"MZ\x90installer"),
        # Ambiguous / Misc
        ("scratch_memo.txt", "Misc", "txt", "Remember to water the plants tomorrow and check mail."),
    ]

    labels_data = []

    for filename, category, filetype, content in samples:
        filepath = base / filename
        labels_data.append({"filename": filename, "expected_category": category})

        if filetype == "pdf":
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 72), content, fontsize=11)
            doc.save(str(filepath))
            doc.close()
        elif filetype in ["txt", "code"]:
            filepath.write_text(content, encoding="utf-8")
        elif filetype == "docx":
            d = docx.Document()
            d.add_heading("Internal Document", level=1)
            d.add_paragraph(content)
            d.save(str(filepath))
        elif filetype == "binary":
            filepath.write_bytes(content)
        elif filetype == "img":
            # Generate a PNG image
            img = Image.new("RGB", (600, 300), color=(240, 244, 248))
            draw = ImageDraw.Draw(img)
            draw.rectangle([20, 20, 580, 280], outline=(100, 150, 200), width=3)
            draw.text((40, 50), "SCREENSHOT PREVIEW", fill=(30, 30, 30))
            draw.text((40, 100), content, fill=(60, 60, 60))
            img.save(filepath)

    with open(labels_file, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["filename", "expected_category"])
        writer.writeheader()
        writer.writerows(labels_data)

    return base


if __name__ == "__main__":
    p = generate_synthetic_dataset()
    print(f"Generated synthetic dataset in {p}")
