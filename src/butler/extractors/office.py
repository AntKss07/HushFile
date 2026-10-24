"""Office documents (DOCX, XLSX, CSV) extractor."""

from __future__ import annotations

import csv
from pathlib import Path

from butler.config import ButlerConfig
from butler.extractors import ExtractionResult


def extract_office(path: Path, config: ButlerConfig) -> ExtractionResult:
    """Extract text from DOCX, XLSX, or CSV."""
    suffix = path.suffix.lower()
    max_chars = config.extract.max_chars

    try:
        if suffix == ".docx":
            import docx
            doc = docx.Document(str(path))
            parts: list[str] = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for t in doc.tables:
                for row in t.rows[:10]:  # Up to 10 rows per table
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        parts.append(row_txt)
            text = "\n".join(parts)[:max_chars]
            return ExtractionResult(
                text=text,
                metadata={"filename": path.name, "extension": ".docx", "paragraphs": len(doc.paragraphs)},
            )

        elif suffix in [".xlsx", ".xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for sheetname in wb.sheetnames[:5]:
                sheet = wb[sheetname]
                parts.append(f"Sheet: {sheetname}")
                row_count = 0
                for row in sheet.iter_rows(values_only=True):
                    if row_count >= 15:  # Headers + top 14 rows
                        break
                    vals = [str(cell) for cell in row if cell is not None]
                    if vals:
                        parts.append(" | ".join(vals))
                    row_count += 1
            wb.close()
            text = "\n".join(parts)[:max_chars]
            return ExtractionResult(
                text=text,
                metadata={"filename": path.name, "extension": suffix, "sheets": wb.sheetnames},
            )

        elif suffix == ".csv":
            parts = []
            # Try utf-8 then latin-1
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    reader = csv.reader(f)
                    for i, row in enumerate(reader):
                        if i >= 20:
                            break
                        if row:
                            parts.append(" | ".join(row))
            except Exception:
                with open(path, "r", encoding="latin-1", errors="replace") as f:
                    reader = csv.reader(f)
                    for i, row in enumerate(reader):
                        if i >= 20:
                            break
                        if row:
                            parts.append(" | ".join(row))

            text = "\n".join(parts)[:max_chars]
            return ExtractionResult(
                text=text,
                metadata={"filename": path.name, "extension": ".csv", "rows_sampled": len(parts)},
            )

        return ExtractionResult(
            text="",
            metadata={"filename": path.name, "extension": suffix},
            warning=f"Unsupported office extension {suffix}",
        )

    except Exception as e:
        return ExtractionResult(
            text="",
            metadata={"filename": path.name, "extension": suffix},
            warning=f"Office extractor error: {e}",
        )
